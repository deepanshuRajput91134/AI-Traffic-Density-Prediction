import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path):
    prs = Presentation()
    # 16:9 Widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    BG_COLOR = RGBColor(15, 23, 42)        # Slate 900
    CARD_BG = RGBColor(30, 41, 59)         # Slate 800
    CARD_BORDER = RGBColor(51, 65, 85)     # Slate 700
    PRIMARY = RGBColor(59, 130, 246)       # Blue 500
    ACCENT = RGBColor(16, 185, 129)        # Emerald 500
    TEXT_MAIN = RGBColor(248, 250, 252)    # Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184)   # Slate 400
    CARD_GREEN = RGBColor(16, 185, 129)
    CARD_YELLOW = RGBColor(245, 158, 11)
    CARD_RED = RGBColor(239, 68, 68)
    CARD_PURPLE = RGBColor(139, 92, 246)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, tag, title):
        # Category Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.4))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = tag.upper()
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = PRIMARY

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_MAIN

    def add_card(slide, left, top, width, height, title, body_bullets, accent_color=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = accent_color if accent_color else CARD_BORDER
        card.line.width = Pt(1.5)

        tx_box = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.2), Inches(width - 0.4), Inches(height - 0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True

        if title:
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(16)
            p.font.bold = True
            p.font.color.rgb = accent_color if accent_color else TEXT_MAIN
            p.space_after = Pt(8)

        for b in body_bullets:
            p = tf.add_paragraph()
            p.text = f"•  {b}"
            p.font.size = Pt(13)
            p.font.color.rgb = TEXT_MUTED
            p.space_after = Pt(5)

    # ═════════════════════════════════════════════════════════════
    # SLIDE 1: TITLE SLIDE
    # ═════════════════════════════════════════════════════════════
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # College Header
    col_box = s1.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.733), Inches(0.8))
    tf_col = col_box.text_frame
    p_col = tf_col.paragraphs[0]
    p_col.text = "IMS ENGINEERING COLLEGE, GHAZIABAD"
    p_col.font.size = Pt(16)
    p_col.font.bold = True
    p_col.font.color.rgb = PRIMARY
    p_col.alignment = PP_ALIGN.CENTER
    p_subcol = tf_col.add_paragraph()
    p_subcol.text = "Department of Computer Science & Engineering (AIML) & CSD"
    p_subcol.font.size = Pt(13)
    p_subcol.font.color.rgb = TEXT_MUTED
    p_subcol.alignment = PP_ALIGN.CENTER

    # Title Card
    t_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.7), Inches(10.333), Inches(2.2))
    t_card.fill.solid()
    t_card.fill.fore_color.rgb = CARD_BG
    t_card.line.color.rgb = PRIMARY
    t_card.line.width = Pt(2)

    tf_t = t_card.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = "AI-Based Traffic Density Prediction\n& Route Optimization"
    p_t.font.size = Pt(28)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_MAIN
    p_t.alignment = PP_ALIGN.CENTER

    p_badge = tf_t.add_paragraph()
    p_badge.text = "Smart City Management System  |  Group No: 5  |  Session 2026-27"
    p_badge.font.size = Pt(13)
    p_badge.font.color.rgb = ACCENT
    p_badge.alignment = PP_ALIGN.CENTER
    p_badge.space_before = Pt(8)

    # Submitted By & Submitted To
    # Submitted By Box
    by_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(4.2), Inches(6.8), Inches(2.6))
    by_card.fill.solid()
    by_card.fill.fore_color.rgb = CARD_BG
    by_card.line.color.rgb = CARD_BORDER
    tf_by = by_card.text_frame
    tf_by.word_wrap = True
    p = tf_by.paragraphs[0]
    p.text = "SUBMITTED BY (TEAM MEMBERS):"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    members = [
        ("Deepanshu Kumar", "2401431530021", "ML Model & Dataset"),
        ("Kaushal Kumar", "2401431530036", "Backend API & Integration"),
        ("Sachin Kashyap", "2401431530051", "Documentation & Presentation"),
        ("Ayush Dwivedi", "2401431530015", "Frontend UI Development")
    ]
    for name, roll, role in members:
        p = tf_by.add_paragraph()
        p.text = f"•  {name}  (Roll No: {roll})  —  {role}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MAIN

    # Submitted To Box
    to_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(4.2), Inches(3.333), Inches(2.6))
    to_card.fill.solid()
    to_card.fill.fore_color.rgb = CARD_BG
    to_card.line.color.rgb = CARD_BORDER
    tf_to = to_card.text_frame
    tf_to.word_wrap = True
    p = tf_to.paragraphs[0]
    p.text = "UNDER GUIDANCE OF:"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    p = tf_to.add_paragraph()
    p.text = "Dr. Nishant Anand"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(8)

    p = tf_to.add_paragraph()
    p.text = "Project Coordinator\nDept of CSE-AIML & CSD\nIMS Engineering College"
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED

    # ═════════════════════════════════════════════════════════════
    # SLIDE 2: PROBLEM STATEMENT & MOTIVATION
    # ═════════════════════════════════════════════════════════════
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Introduction", "Problem Statement & Urban Traffic Challenges")

    add_card(s2, 0.8, 1.8, 3.6, 5.0, "Urban Congestion Crisis", [
        "Rapid urbanization causes severe bottlenecks in metropolitan road networks.",
        "Traffic congestion costs billions of dollars annually in wasted fuel and time.",
        "Commuters spend 100+ unproductive hours stranded in traffic jams each year."
    ], CARD_RED)

    add_card(s2, 4.8, 1.8, 3.6, 5.0, "Flaws in Existing Systems", [
        "Fixed-interval traffic signal timers fail during unexpected rush hours.",
        "Static navigation algorithms don't account for dynamic AI-predicted road load.",
        "Lack of centralized data persistence for municipal traffic analysis."
    ], CARD_YELLOW)

    add_card(s2, 8.8, 1.8, 3.6, 5.0, "Proposed AI Solution", [
        "Predictive Machine Learning models that forecast congestion before it occurs.",
        "Traffic-weighted Dijkstra graph algorithm for optimal route dispatch.",
        "End-to-end Smart City dashboard with live REST API and SQLite history."
    ], CARD_GREEN)

    # ═════════════════════════════════════════════════════════════
    # SLIDE 3: OBJECTIVES & SCOPE
    # ═════════════════════════════════════════════════════════════
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Goals & Deliverables", "Project Objectives & Scope of Work")

    add_card(s3, 0.8, 1.8, 5.6, 2.4, "1. Predictive Intelligence", [
        "Train & benchmark multi-class ML classifiers (Low / Medium / High).",
        "Evaluate with 5-fold Stratified Cross-Validation on urban dataset.",
        "Incorporate temporal, weather, and vehicular density features."
    ], PRIMARY)

    add_card(s3, 6.8, 1.8, 5.6, 2.4, "2. Dynamic Route Optimization", [
        "Implement Dijkstra's algorithm on a 6-node city road network.",
        "Dynamically penalize edge weights based on real-time traffic density.",
        "Provide primary shortest route alongside viable alternate routes."
    ], ACCENT)

    add_card(s3, 0.8, 4.5, 5.6, 2.4, "3. Scalable Web Architecture", [
        "High-performance REST API using Python FastAPI & Pydantic.",
        "SQLite database persistence for historical queries and analytics.",
        "CORS middleware for seamless frontend communication."
    ], CARD_PURPLE)

    add_card(s3, 6.8, 4.5, 5.6, 2.4, "4. Interactive Dashboard & CV", [
        "Real-time React 18 + Vite dashboard with tabs for Prediction & Routing.",
        "Computer Vision module with OpenCV for camera-assisted vehicle counting.",
        "Containerized with Docker & Docker Compose for instant deployment."
    ], CARD_YELLOW)

    # ═════════════════════════════════════════════════════════════
    # SLIDE 4: SYSTEM ARCHITECTURE
    # ═════════════════════════════════════════════════════════════
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Engineering Design", "End-to-End System Architecture")

    # Layer 1: Client
    add_card(s4, 0.8, 1.8, 11.6, 1.3, "PRESENTATION LAYER  (Frontend)", [
        "React 18  •  Vite Dev Server  •  Interactive Dashboard  •  Prediction Form  •  Routing Map  •  Analytics & Vision Tabs"
    ], PRIMARY)

    # Layer 2: API
    add_card(s4, 0.8, 3.4, 11.6, 1.3, "APPLICATION & API LAYER  (FastAPI Backend)", [
        "FastAPI Service (Port 8000)  •  Modular Routers: /traffic, /routes, /analytics, /vision  •  Pydantic Validation  •  CORS"
    ], ACCENT)

    # Layer 3: Services
    add_card(s4, 0.8, 5.0, 3.6, 1.8, "ML INFERENCE ENGINE", [
        "Random Forest Classifier",
        "Joblib serialized pipeline",
        "5-fold CV champion model"
    ], CARD_YELLOW)

    add_card(s4, 4.8, 5.0, 3.6, 1.8, "GRAPH ROUTING SERVICE", [
        "Dijkstra's shortest path",
        "6 nodes & 9 weighted edges",
        "Traffic congestion penalty"
    ], CARD_PURPLE)

    add_card(s4, 8.8, 5.0, 3.6, 1.8, "PERSISTENCE & VISION", [
        "SQLite Database (traffic.db)",
        "traffic_log & route_log tables",
        "OpenCV vehicle detector stub"
    ], CARD_GREEN)

    # ═════════════════════════════════════════════════════════════
    # SLIDE 5: ML MODEL DEVELOPMENT & BENCHMARKING
    # ═════════════════════════════════════════════════════════════
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Machine Learning", "Model Selection & Cross-Validation Benchmarking")

    add_card(s5, 0.8, 1.8, 5.6, 5.0, "Model Comparison & Selection", [
        "Trained and evaluated 4 distinct classifiers on curated traffic dataset:",
        "🥇 Random Forest Classifier (Selected Champion)",
        "    • Superior non-linear feature handling and ensemble stability.",
        "    • Robust against overfitting via bagging and tree averaging.",
        "🥈 Gradient Boosting (Runner-up) — High accuracy, slightly slower.",
        "🥉 Decision Tree (Baseline) — Prone to high variance.",
        "4️⃣ Logistic Regression — Evaluated as linear baseline.",
        "All models evaluated using 5-Fold Stratified Cross-Validation."
    ], CARD_GREEN)

    add_card(s5, 6.8, 1.8, 5.6, 5.0, "Features & Prediction Output", [
        "Key Input Features Utilized:",
        "  • Hour of Day (0–23) — Captures peak morning & evening rush hours.",
        "  • Day of Week (0–6) — Differentiates weekdays from weekends.",
        "  • Vehicle Count — Quantitative density input from sensors/cameras.",
        "  • Weather Condition — Rain, Fog, Clear, Snow impact on congestion.",
        "  • Road Type — Highway, Arterial, Local Street capacity parameters.",
        "Target Classes Output:",
        "  🟢 Low Density  |  🟡 Medium Density  |  🔴 High Density",
        "Artifacts Generated: traffic_model.pkl, label_encoder.pkl"
    ], PRIMARY)

    # ═════════════════════════════════════════════════════════════
    # SLIDE 6: DIJKSTRA GRAPH ROUTING ALGORITHM
    # ═════════════════════════════════════════════════════════════
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Algorithms", "Traffic-Aware Dijkstra Graph Route Optimization")

    add_card(s6, 0.8, 1.8, 5.6, 5.0, "Smart City Graph Topology", [
        "Simulated 6 major urban junction hubs:",
        "  1. City Center (Central Business District)",
        "  2. Tech Park (IT & Corporate Corridor)",
        "  3. North Junction (Residential Connector)",
        "  4. South Interchange (Highway Entry Point)",
        "  5. East Corridor (Commercial Suburb)",
        "  6. Airport (Major Transit Hub)",
        "9 Bidirectional Weighted Edges connecting hubs.",
        "Base edge weights represent standard travel distance & transit time in minutes."
    ], PRIMARY)

    add_card(s6, 6.8, 1.8, 5.6, 5.0, "Dynamic Traffic Penalty Formula", [
        "Standard Dijkstra minimizes base weight:  W(u, v)",
        "Dynamic Traffic Penalty applied in our system:",
        "  • Low Density:    W_eff = W_base × 1.0 (No delay)",
        "  • Medium Density: W_eff = W_base × 1.4 (+40% transit penalty)",
        "  • High Density:   W_eff = W_base × 2.0 (+100% heavy jam penalty)",
        "Algorithm outputs:",
        "  ✓ Primary optimal path with lowest effective travel cost.",
        "  ✓ Alternative backup paths for rerouting during road blockages.",
        "All routing decisions automatically recorded to SQLite route_log."
    ], ACCENT)

    # ═════════════════════════════════════════════════════════════
    # SLIDE 7: BACKEND API & PERSISTENCE
    # ═════════════════════════════════════════════════════════════
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Backend Infrastructure", "FastAPI REST Endpoints & SQLite Persistence")

    add_card(s7, 0.8, 1.8, 3.6, 5.0, "Core API Endpoints", [
        "GET /health",
        "  • Real-time status of DB & ML model.",
        "POST /traffic/predict",
        "  • Takes input features & returns density.",
        "GET /routes/smart-predict",
        "  • Dijkstra optimal path calculation.",
        "GET /analytics/summary",
        "  • Aggregated counts & logs.",
        "GET /vision/vision-demo",
        "  • Vehicle detector simulation."
    ], PRIMARY)

    add_card(s7, 4.8, 1.8, 3.6, 5.0, "Database Architecture", [
        "SQLite Engine (traffic.db):",
        "• traffic_log table:",
        "   - id, timestamp, hour, day, count, weather, predicted_density.",
        "• route_log table:",
        "   - start_point, end_point, chosen_route, travel_time, traffic_level.",
        "• analytics table:",
        "   - Precomputed aggregation metrics.",
        "Auto-initialized on FastAPI startup."
    ], CARD_PURPLE)

    add_card(s7, 8.8, 1.8, 3.6, 5.0, "Production Resilience", [
        "• Lazy model loader prevents crashes if artifact is temporarily missing.",
        "• Pydantic schema validation blocks malformed requests.",
        "• Strict CORS headers enable dev and production host communication.",
        "• Auto-generated interactive Swagger API documentation at /docs."
    ], CARD_GREEN)

    # ═════════════════════════════════════════════════════════════
    # SLIDE 8: REACT DASHBOARD & FRONTEND UI
    # ═════════════════════════════════════════════════════════════
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Frontend UI", "Modern React 18 & Vite Interactive Dashboard")

    add_card(s8, 0.8, 1.8, 5.6, 2.4, "Interactive Tabbed Navigation", [
        "Clean, responsive tabbed interface built with React 18 & Vite.",
        "Tabs: Traffic Prediction  |  Route Optimizer  |  Analytics  |  Vision Demo.",
        "Real-time system health pill (Backend, Database, ML model status)."
    ], PRIMARY)

    add_card(s8, 6.8, 1.8, 5.6, 2.4, "Prediction & Simulation", [
        "Interactive sliders and dropdowns for Hour, Weather, and Vehicle Count.",
        "Instant visual traffic badge (Low / Medium / High) with congestion ratio.",
        "Async API calls with fallback error handling."
    ], ACCENT)

    add_card(s8, 0.8, 4.5, 5.6, 2.4, "Visual Route Optimization", [
        "Origin & Destination selectors across all 6 smart city nodes.",
        "Visual path display showing optimal route sequence and estimated time.",
        "Traffic penalty impact highlighted to explain AI recommendation."
    ], CARD_YELLOW)

    add_card(s8, 6.8, 4.5, 5.6, 2.4, "Analytics & Computer Vision", [
        "Real-time analytics summary table populated directly from SQLite logs.",
        "Computer vision frame simulation showing vehicle count and density.",
        "Lightweight JavaScript fetch API wrapper — zero third-party client bloat."
    ], CARD_PURPLE)

    # ═════════════════════════════════════════════════════════════
    # SLIDE 9: COMPUTER VISION MODULE
    # ═════════════════════════════════════════════════════════════
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Computer Vision", "Camera-Assisted Vehicle Detection with OpenCV")

    add_card(s9, 0.8, 1.8, 5.6, 5.0, "Computer Vision Pipeline", [
        "Purpose: Bridge real-world traffic cameras with the AI prediction engine.",
        "Core Processing Steps:",
        "  1. Video Stream / Frame Acquisition (CCTV traffic feeds).",
        "  2. Grayscale conversion & Gaussian Blur filtering.",
        "  3. Background subtraction to isolate moving vehicular silhouettes.",
        "  4. Contour detection & bounding box thresholding.",
        "  5. Automated vehicular counting within regions of interest (ROI).",
        "Vehicle count directly feeds as input feature to ML prediction router."
    ], PRIMARY)

    add_card(s9, 6.8, 1.8, 5.6, 5.0, "Architecture & Future Vision Upgrades", [
        "Current Implementation:",
        "  • OpenCV 5.0 detector module in computer-vision/vehicle_detector.py.",
        "  • REST endpoint GET /vision/vision-demo for simulation testing.",
        "Future Enhancements for Final Year:",
        "  • Integrate Ultralytics YOLOv8 for vehicle classification (car, bus, truck, bike).",
        "  • Real-time RTSP camera stream processing using edge AI (NVIDIA Jetson).",
        "  • Automatic License Plate Recognition (ALPR) for traffic law enforcement."
    ], CARD_GREEN)

    # ═════════════════════════════════════════════════════════════
    # SLIDE 10: TEAM WORK DISTRIBUTION
    # ═════════════════════════════════════════════════════════════
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Collaboration", "Team Work Distribution & Contribution")

    add_card(s10, 0.8, 1.8, 2.7, 5.0, "Deepanshu Kumar", [
        "Roll No: 2401431530021",
        "Role: ML & Datasets",
        "• Data curation & feature engineering.",
        "• Random Forest training.",
        "• 5-Fold Stratified Cross-Validation.",
        "• Model evaluation & metadata generation."
    ], CARD_GREEN)

    add_card(s10, 3.8, 1.8, 2.7, 5.0, "Kaushal Kumar", [
        "Roll No: 2401431530036",
        "Role: Backend & API",
        "• FastAPI app architecture.",
        "• Dijkstra routing service.",
        "• SQLite schema & persistence helpers.",
        "• REST API endpoints & CORS setup."
    ], CARD_YELLOW)

    add_card(s10, 6.8, 1.8, 2.7, 5.0, "Ayush Dwivedi", [
        "Roll No: 2401431530015",
        "Role: Frontend Dev",
        "• React 18 + Vite dashboard.",
        "• Tabbed UI navigation.",
        "• Prediction & Routing forms.",
        "• API service integration layer."
    ], CARD_RED)

    add_card(s10, 9.8, 1.8, 2.7, 5.0, "Sachin Kashyap", [
        "Roll No: 2401431530051",
        "Role: Documentation",
        "• Synopsis & Project Report.",
        "• Presentation Slide Deck.",
        "• Docker Compose setup.",
        "• Deployment documentation."
    ], CARD_PURPLE)

    # ═════════════════════════════════════════════════════════════
    # SLIDE 11: TESTING, DEVOPS & RESULTS
    # ═════════════════════════════════════════════════════════════
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Validation & DevOps", "Automated Testing, Dockerization & Live Metrics")

    add_card(s11, 0.8, 1.8, 5.6, 5.0, "Automated Test Suite (13/13 Passed)", [
        "Comprehensive test suite covering all project subsystems:",
        "✓ test_api.py: Verified /health, /predict, /smart-predict, /status.",
        "✓ test_ml.py: Validated inference shape, classes, and latency.",
        "✓ test_vision.py: Confirmed OpenCV vehicle counting stub output.",
        "✓ test_analytics.py: Tested SQLite CRUD and aggregation queries.",
        "All 13 unit & integration tests pass with 0 regressions.",
        "Fast response time: < 25ms average inference latency on CPU."
    ], CARD_GREEN)

    add_card(s11, 6.8, 1.8, 5.6, 5.0, "DevOps & Cloud Readiness", [
        "🐳 Multi-Stage Dockerfile:",
        "  • Stage 1: Builds React 18 frontend with Node.js 20 Alpine.",
        "  • Stage 2: Packages Python 3.12 backend + OpenCV libraries.",
        "🔧 Docker Compose (docker-compose.yml):",
        "  • One-command startup for both backend & frontend containers.",
        "  • Persistent volume mounting for traffic.db.",
        "🐙 GitHub Repository:",
        "  • Clean Git commit history, .gitignore, and .env configuration.",
        "  • 3 team collaborators added with full write access."
    ], PRIMARY)

    # ═════════════════════════════════════════════════════════════
    # SLIDE 12: CONCLUSION & FUTURE SCOPE
    # ═════════════════════════════════════════════════════════════
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Summary & Roadmap", "Conclusion & Future Scope")

    add_card(s12, 0.8, 1.8, 5.6, 5.0, "Conclusion", [
        "Successfully engineered a complete end-to-end prototype for Smart City Traffic Management.",
        "Demonstrated synergy between ML prediction, graph optimization, and real-time visualization.",
        "Key Achievements:",
        "  ✓ Trained champion Random Forest model with 5-fold CV.",
        "  ✓ Dynamic Dijkstra routing with traffic-weighted penalties.",
        "  ✓ Full-stack implementation (FastAPI + React 18 + SQLite).",
        "  ✓ 13/13 automated tests passing.",
        "  ✓ Docker containerized and cloud-deployment ready."
    ], CARD_GREEN)

    add_card(s12, 6.8, 1.8, 5.6, 5.0, "Future Scope for Final Year", [
        "1. Real-Time IoT Sensor Integration: Ingest live ultrasonic and inductive loop traffic sensors.",
        "2. GPS & GIS Mapping: Integrate Mapbox / OpenStreetMap with live geolocation polyline routes.",
        "3. Deep Learning Upgrade: Deploy YOLOv8 for multi-class vehicle detection on camera streams.",
        "4. Smart Traffic Signal Automation: Dynamic signal timer adjustment based on queue length.",
        "5. Cloud Deployment: Deploy on AWS EC2 / Render with automated CI/CD pipelines."
    ], PRIMARY)

    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "AI_Traffic_Management_Presentation.pptx"
    create_deck(out)
